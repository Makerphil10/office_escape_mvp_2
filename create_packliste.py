from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter

wb = Workbook()

# ── Farben ──────────────────────────────────────────────────────────────
DARK_BLUE   = "1E3A5F"
WHITE       = "FFFFFF"
LIGHT_GREY  = "F1F5F9"
STRIPE      = "FAFAFA"
BORDER_COL  = "E2E8F0"

COL_PH  = "DBEAFE"   # blau
COL_ST  = "FCE7F3"   # rosa
COL_EM  = "D1FAE5"   # gruen
COL_AN  = "FEF3C7"   # gelb
COL_AL  = "EDE9FE"   # lila
COL_GEM = "F8F9FF"   # gemeinsam

# hellere Varianten fuer Datenzellen
COL_PH2 = "EFF6FF"
COL_ST2 = "FDF2F8"
COL_EM2 = "ECFDF5"
COL_AN2 = "FFFBEB"
COL_AL2 = "F5F3FF"

def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def thin_border():
    s = Side(style='thin', color=BORDER_COL)
    return Border(left=s, right=s, top=s, bottom=s)

def header_font(bold=True, color=WHITE, size=11):
    return Font(name="Segoe UI", bold=bold, color=color, size=size)

def cell_font(bold=False, color="1A202C", size=10):
    return Font(name="Segoe UI", bold=bold, color=color, size=size)

def center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def left():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

# ── Sheet-Hilfsfunktionen ────────────────────────────────────────────────

def set_col_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

def add_title_row(ws, row, title, icon=""):
    ws.row_dimensions[row].height = 22
    for c in range(1, 8):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill(DARK_BLUE)
        cell.border = thin_border()
    ws.cell(row=row, column=1).value = f"  {icon}  {title}"
    ws.cell(row=row, column=1).font = Font(name="Segoe UI", bold=True, color=WHITE, size=12)
    ws.cell(row=row, column=1).alignment = left()
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)

def add_col_headers(ws, row, labels):
    ws.row_dimensions[row].height = 36
    bg_colors = [LIGHT_GREY, COL_PH, COL_ST, COL_EM, COL_AN, COL_AL, LIGHT_GREY]
    for c, (label, bg) in enumerate(zip(labels, bg_colors), start=1):
        cell = ws.cell(row=row, column=c)
        cell.value = label
        cell.fill = fill(bg)
        cell.font = Font(name="Segoe UI", bold=True, color=DARK_BLUE, size=9)
        cell.alignment = center()
        cell.border = thin_border()

def add_data_row(ws, row, values, stripe=False, gemeinsam=False):
    ws.row_dimensions[row].height = 18
    bg = [STRIPE if stripe else WHITE, COL_PH2, COL_ST2, COL_EM2, COL_AN2, COL_AL2, LIGHT_GREY]
    if gemeinsam:
        bg = [COL_GEM] * 7

    for c, (val, bgc) in enumerate(zip(values, bg), start=1):
        cell = ws.cell(row=row, column=c)
        cell.value = val
        cell.fill = fill(bgc)
        cell.border = thin_border()
        if c == 1:
            cell.font = Font(name="Segoe UI", bold=False, color="1A202C", size=10)
            cell.alignment = left()
        elif c == 7:
            cell.alignment = center()
        else:
            cell.font = Font(name="Segoe UI", bold=True, color=DARK_BLUE, size=11)
            cell.alignment = center()

def add_gemeinsam_row(ws, row, item, value, stripe=False):
    ws.row_dimensions[row].height = 18
    bg = STRIPE if stripe else COL_GEM
    for c in range(1, 8):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill(bg)
        cell.border = thin_border()

    ws.cell(row=row, column=1).value = item
    ws.cell(row=row, column=1).font = Font(name="Segoe UI", size=10, color="1A202C")
    ws.cell(row=row, column=1).alignment = left()

    ws.cell(row=row, column=2).value = value
    ws.cell(row=row, column=2).font = Font(name="Segoe UI", bold=True, color=DARK_BLUE, size=10)
    ws.cell(row=row, column=2).alignment = center()
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)

    # Checkbox-Spalte
    ws.cell(row=row, column=7).value = ""
    ws.cell(row=row, column=7).alignment = center()

def add_empty_row(ws, row, height=8):
    ws.row_dimensions[row].height = height
    for c in range(1, 8):
        ws.cell(row=row, column=c).value = ""

# ── SHEET 1: KLEIDUNG ────────────────────────────────────────────────────
ws = wb.active
ws.title = "Kleidung"

set_col_widths(ws, {
    "A": 38, "B": 13, "C": 13, "D": 13, "E": 13, "F": 13, "G": 8
})

# Haupttitel
ws.row_dimensions[1].height = 40
for c in range(1, 8):
    ws.cell(row=1, column=c).fill = fill(DARK_BLUE)
    ws.cell(row=1, column=c).border = thin_border()
ws.cell(row=1, column=1).value = "PACKLISTE  ·  FAMILIE HECK"
ws.cell(row=1, column=1).font = Font(name="Segoe UI", bold=True, color=WHITE, size=16)
ws.cell(row=1, column=1).alignment = left()
ws.merge_cells("A1:G1")

# Reiseinfo
ws.row_dimensions[2].height = 20
ws.cell(row=2, column=1).value = "Reiseziel: ___________________     Zeitraum: ___________________     Anzahl Tage: ______"
ws.merge_cells("A2:G2")

add_empty_row(ws, 3, 6)

# ── Kleidung ──
r = 4
add_title_row(ws, r, "KLEIDUNG", "👕"); r += 1
add_col_headers(ws, r, ["Item", "Philipp", "Steffi", "Emil (10)", "Anton (7)", "Alma (2)", "✓"]); r += 1

kleidung = [
    ("T-Shirts / Tops",                    7,  7,  10, 10, 10),
    ("Kurze Hosen / Shorts",               4,  4,   7,  7,  7),
    ("Lange Hosen / Leggings",             2,  2,   2,  2,  3),
    ("Kleider / Röcke  (Steffi & Alma)",   "—", 3, "—","—",  5),
    ("Longsleeves / Langarmshirts",        1,  1,   3,  3,  4),
    ("Pullover / Strickjacke",             1,  2,   2,  2,  3),
    ("Leichte Jacke",                      1,  1,   1,  1,  1),
    ("Unterwäsche",                       10, 10,  10, 10, "—"),
    ("Socken",                            10, 10,  10, 10, 10),
    ("Strumpfhosen  (Steffi & Alma)",     "—","2","—","—",  4),
    ("Schlafanzug / Schlafi",              1,  1,   2,  2,  3),
    ("Bodies kurz  (Alma)",               "—","—","—","—", 10),
    ("Bodies lang  (Alma)",               "—","—","—","—",  9),
    ("Jumpsuits  (Alma)",                 "—","—","—","—",  2),
    ("Windeln normal  (Alma)",            "—","—","—","—", 40),
    ("Lätzchen / Sabberlätzchen  (Alma)", "—","—","—","—",  5),
]

for i, (item, ph, st, em, an, al) in enumerate(kleidung):
    add_data_row(ws, r, [item, ph, st, em, an, al, ""], stripe=(i % 2 == 1))
    r += 1

add_empty_row(ws, r, 6); r += 1

# ── Badesachen ──
add_title_row(ws, r, "BADESACHEN", "🏊"); r += 1
add_col_headers(ws, r, ["Item", "Philipp", "Steffi", "Emil (10)", "Anton (7)", "Alma (2)", "✓"]); r += 1

badesachen = [
    ("Badehose / Badeanzug",       2, 2, 2,   2,  2),
    ("UV-Shirt / Badeshirt",      "—","—", 2,  2,  2),
    ("Schwimmwindeln  (Alma)",    "—","—","—","—", 15),
]
for i, (item, ph, st, em, an, al) in enumerate(badesachen):
    add_data_row(ws, r, [item, ph, st, em, an, al, ""], stripe=(i % 2 == 1))
    r += 1

add_empty_row(ws, r, 6); r += 1

# ── Schuhe & Accessoires ──
add_title_row(ws, r, "SCHUHE & ACCESSOIRES", "👟"); r += 1
add_col_headers(ws, r, ["Item", "Philipp", "Steffi", "Emil (10)", "Anton (7)", "Alma (2)", "✓"]); r += 1

schuhe = [
    ("Sandalen / Flipflops / Schlappen",  1, 1, 1, 1, 1),
    ("Sneaker / geschlossene Schuhe",     1, 1, 1, 1, 1),
    ("Wasserschuhe",                     "—","—",1, 1, 1),
    ("Schickere Schuhe  (optional, Abend)", 1, 1,"—","—","—"),
    ("Sonnenhut / Käppi",                 1, 1, 1, 1, 1),
    ("Sonnenbrille",                      1, 1, 1, 1,"—"),
    ("Haargummi  (Steffi & Alma)",       "—","✓","—","—","✓"),
]
for i, (item, ph, st, em, an, al) in enumerate(schuhe):
    add_data_row(ws, r, [item, ph, st, em, an, al, ""], stripe=(i % 2 == 1))
    r += 1

add_empty_row(ws, r, 6); r += 1

# ── Abendoutfit ──
add_title_row(ws, r, "ABENDOUTFIT / BESONDERER ANLASS  (optional)", "👔"); r += 1
add_col_headers(ws, r, ["Item", "Philipp", "Steffi", "Emil (10)", "Anton (7)", "Alma (2)", "✓"]); r += 1

abend = [
    ("Hemd / festliches Kleid",     1, 1, 1, 1, 1),
    ("Festliche Hose / Rock",       1,"—",1, 1,"—"),
    ("Fliege / Krawatte",          "—","—",1, 1,"—"),
]
for i, (item, ph, st, em, an, al) in enumerate(abend):
    add_data_row(ws, r, [item, ph, st, em, an, al, ""], stripe=(i % 2 == 1))
    r += 1

ws.freeze_panes = "A6"
ws.print_title_rows = "1:5"


# ── SHEET 2: AUSRÜSTUNG & EXTRAS ────────────────────────────────────────
ws2 = wb.create_sheet("Ausruestung & Extras")
set_col_widths(ws2, {"A": 46, "B": 18, "C": 8})

def add_title2(ws, row, title, icon=""):
    ws.row_dimensions[row].height = 22
    for c in range(1, 4):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill(DARK_BLUE)
        cell.border = thin_border()
    ws.cell(row=row, column=1).value = f"  {icon}  {title}"
    ws.cell(row=row, column=1).font = Font(name="Segoe UI", bold=True, color=WHITE, size=12)
    ws.cell(row=row, column=1).alignment = left()
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)

def add_header2(ws, row):
    ws.row_dimensions[row].height = 28
    for c, (lbl, bg) in enumerate(zip(["Item", "Anzahl / Hinweis", "✓"],
                                       [LIGHT_GREY, LIGHT_GREY, LIGHT_GREY]), start=1):
        cell = ws.cell(row=row, column=c)
        cell.value = lbl
        cell.fill = fill(bg)
        cell.font = Font(name="Segoe UI", bold=True, color=DARK_BLUE, size=9)
        cell.alignment = center()
        cell.border = thin_border()

def add_row2(ws, row, item, value, stripe=False):
    ws.row_dimensions[row].height = 18
    bg = STRIPE if stripe else WHITE
    ws.cell(row=row, column=1).value = item
    ws.cell(row=row, column=1).fill = fill(COL_GEM if not stripe else STRIPE)
    ws.cell(row=row, column=1).font = Font(name="Segoe UI", size=10, color="1A202C")
    ws.cell(row=row, column=1).alignment = left()
    ws.cell(row=row, column=1).border = thin_border()

    ws.cell(row=row, column=2).value = value
    ws.cell(row=row, column=2).fill = fill(COL_GEM if not stripe else STRIPE)
    ws.cell(row=row, column=2).font = Font(name="Segoe UI", bold=True, color=DARK_BLUE, size=10)
    ws.cell(row=row, column=2).alignment = center()
    ws.cell(row=row, column=2).border = thin_border()

    ws.cell(row=row, column=3).value = ""
    ws.cell(row=row, column=3).fill = fill(LIGHT_GREY)
    ws.cell(row=row, column=3).alignment = center()
    ws.cell(row=row, column=3).border = thin_border()

def add_empty2(ws, row):
    ws.row_dimensions[row].height = 8
    for c in range(1, 4):
        ws.cell(row=row, column=c).value = ""

ws2.row_dimensions[1].height = 40
for c in range(1, 4):
    ws2.cell(row=1, column=c).fill = fill(DARK_BLUE)
    ws2.cell(row=1, column=c).border = thin_border()
ws2.cell(row=1, column=1).value = "PACKLISTE  ·  FAMILIE HECK  ·  Ausrüstung & Extras"
ws2.cell(row=1, column=1).font = Font(name="Segoe UI", bold=True, color=WHITE, size=14)
ws2.cell(row=1, column=1).alignment = left()
ws2.merge_cells("A1:C1")

add_empty2(ws2, 2)

r2 = 3

add_title2(ws2, r2, "STRAND & POOL", "🏗️"); r2 += 1
add_header2(ws2, r2); r2 += 1
strand = [
    ("Badehandtücher / Strandtücher",                " 5 Stück"),
    ("Strandtasche / Rucksack",                      "1"),
    ("Schwimmbrillen",                               "3 (Emil, Anton, Philipp)"),
    ("Schnorchelbrillen",                            "2"),
    ("Schwimmflügel / Schwimmhilfe  (Alma)",         "1"),
    ("UV-Zelt / Strandmuschel",                      "1"),
    ("Strandspielzeug  (Schaufel, Eimer, Formen)",   "✓"),
    ("Wasserschuhe  (Emil, Anton, Alma)",             "→ siehe Kleidung"),
    ("Babydecke  (Alma)",                            "1"),
]
for i, (item, val) in enumerate(strand):
    add_row2(ws2, r2, item, val, stripe=(i % 2 == 1)); r2 += 1

add_empty2(ws2, r2); r2 += 1

add_title2(ws2, r2, "PFLEGE & GESUNDHEIT", "💈"); r2 += 1
add_header2(ws2, r2); r2 += 1
pflege = [
    ("Sonnencreme LSF 30–50",                                         "mehrere Flaschen!"),
    ("After Sun / Aloe Vera",                                          "✓"),
    ("Feuchttücher  (Alma)",                                          "2 Pakete"),
    ("Pipi-Unterlagen  (Anton & Alma, für nachts)",                   "6 + 6"),
    ("Reiseapotheke  (Fieberthermometer, Medikamente, Pflaster, Mückenschutz)", "✓"),
    ("Zahnbørsten (5 Stück) + Zahnpasta",                            "✓"),
    ("Shampoo / Duschgel",                                            "✓"),
    ("Nagelschere",                                                   "✓"),
    ("Nachtlicht  (Emil)",                                            "✓"),
    ("Babybrei / Quetschies  (Alma)",                                 "✓"),
]
for i, (item, val) in enumerate(pflege):
    add_row2(ws2, r2, item, val, stripe=(i % 2 == 1)); r2 += 1

add_empty2(ws2, r2); r2 += 1

add_title2(ws2, r2, "DOKUMENTE & GELD", "📄"); r2 += 1
add_header2(ws2, r2); r2 += 1
docs = [
    ("Reisepässe (5 Stück) + Kopien",         "✓"),
    ("Flugtickets / Buchungsbestätigung",      "✓"),
    ("Versicherungskarten  (Kranken, Reise)",   "✓"),
    ("Kreditkarte",                             "✓"),
    ("Bargeld  (Euro + Lokalwährung)",          "✓"),
    ("Führerschein  (Falls Mietwagen)",        "✓"),
]
for i, (item, val) in enumerate(docs):
    add_row2(ws2, r2, item, val, stripe=(i % 2 == 1)); r2 += 1

add_empty2(ws2, r2); r2 += 1

add_title2(ws2, r2, "ELEKTRONIK", "📱"); r2 += 1
add_header2(ws2, r2); r2 += 1
elektronik = [
    ("Smartphones + Ladekabel (2×)",      "✓"),
    ("Powerbank",                          "✓"),
    ("Reiseadapter  (falls nötig)",        "✓"),
    ("Kolfhörer",                          "✓"),
    ("Tablet / Switch + Spiele",           "✓"),
    ("E-Book-Reader / Bøcher",              "✓"),
    ("Kamera",                              "✓"),
]
for i, (item, val) in enumerate(elektronik):
    add_row2(ws2, r2, item, val, stripe=(i % 2 == 1)); r2 += 1

add_empty2(ws2, r2); r2 += 1

add_title2(ws2, r2, "UNTERHALTUNG & EXTRAS", "🎮"); r2 += 1
add_header2(ws2, r2); r2 += 1
extras = [
    ("Lieblingskuscheltier  (Emil, Anton, Alma)",    "3"),
    ("Reise-Spiele / Karten",                        "✓"),
    ("Malzeug / Stickerhefte",                       "✓"),
    ("Trinkflaschen (5×)",                           "✓"),
    ("Snacks  (Flug & Ausfløge)",                    "✓"),
    ("Buggy  (Alma)",                                "✓"),
    ("Babytrage  (Alma)",                            "✓"),
]
for i, (item, val) in enumerate(extras):
    add_row2(ws2, r2, item, val, stripe=(i % 2 == 1)); r2 += 1

ws2.freeze_panes = "A3"

wb.save("/tmp/Packliste_Familie_Heck.xlsx")
print("Fertig!")

